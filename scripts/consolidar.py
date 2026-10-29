"""
Consolida as tabelas da PAM/IBGE (uma planilha por UF) nos arquivos usados pelo painel.

Uso:
    pip install openpyxl
    python scripts/consolidar.py CAMINHO/DA/PASTA/COM/AS/TABELAS

A pasta deve conter os arquivos baixados do SIDRA, com os nomes no padrão
"Tabela_3_<n>_-_<UF>.xlsx" (lavouras temporárias, Tabela 1612),
"Tabela_6_<n>_-_<UF>.xlsx" (lavouras permanentes, Tabela 1613) e
"Tabela_2_<n>_-_<UF>.xlsx" (efetivo dos rebanhos, Tabela 3939).
Para uma UF sem arquivo Tabela_3_*, as lavouras temporárias podem vir da Tabela 5457
baixada por município (qualquer outro .xlsx da pasta, com área, quantidade e valor da
produção, em uma ou mais planilhas), como foi feito para Minas Gerais.
Gera dados/pam_2025.json e dados/ppm_2025.json (lidos pelo painel)
e dados/pam_2025.csv e dados/ppm_2025.csv (download).
O crédito rural vem do arquivo MATRIZ_CR_REGIAO*.csv baixado da Matriz de Dados do Crédito
Rural do Banco Central, colocado na mesma pasta (gera dados/credito_2025.json e .csv).
Os produtos financiados vêm do arquivo MATRIZ_CR_PRODUTO_REGIAO*.csv (base por estado),
que gera dados/credito_produtos_2025.json e .csv.
A produção de origem animal (Tabela 74) e a aquicultura (Tabela 3940), por estado, vêm de
planilhas .xlsx do SIDRA na mesma pasta (gera dados/pecuaria_produtos_uf_2025.csv).
O CAR vem de duas planilhas exportadas do Painel de Regularização Ambiental (número de
cadastros e área cadastrada), na mesma pasta (gera dados/car_2025.json e .csv).
Para outro ano, ajuste ANO abaixo e os nomes dos arquivos em index.html.
"""
import csv, glob, json, os, re, sys, unicodedata
from collections import defaultdict
import openpyxl

ANO = 2025
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UFMAP = {'AC':12,'AL':27,'AP':16,'AM':13,'BA':29,'CE':23,'DF':53,'ES':32,'GO':52,'MA':21,'MT':51,'MS':50,'MG':31,'PA':15,'PB':25,'PR':41,'PE':26,'PI':22,'RJ':33,'RN':24,'RS':43,'RO':11,'RR':14,'SC':42,'SP':35,'SE':28,'TO':17}
UFSIG = {v: k for k, v in UFMAP.items()}
# Grafias das tabelas do IBGE que diferem da lista oficial de municípios.
CORRECOES = {('BA','Muquém do São Francisco'):'Muquém de São Francisco',('BA','Santa Terezinha'):'Santa Teresinha',
 ('MG','Barão do Monte Alto'):'Barão de Monte Alto',('MG','Dona Euzébia'):'Dona Eusébia',('MG','São Tomé das Letras'):'São Thomé das Letras',
 ('MT','Santo Antônio de Leverger'):'Santo Antônio do Leverger',('PE','São Caitano'):'São Caetano',('RN','Arez'):'Arês',('RN','Assú'):'Açu',
 ('RN','Campo Grande'):'Augusto Severo (Campo Grande)',('RN','Januário Cicco'):'Januário Cicco (Boa Saúde)',('RR','São Luiz do Anauá'):'São Luiz',
 ('SE','Amparo do São Francisco'):'Amparo de São Francisco',('SE','Graccho Cardoso'):'Gracho Cardoso',('SP','Florínea'):'Florínia',('TO','Tabocão'):'Fortaleza do Tabocão'}
ROTULOS = {'Café (em grão) Total':'Café (em grão) — total','Café (em grão) Arábica':'Café (em grão) — arábica','Café (em grão) Canephora':'Café (em grão) — canephora'}
UNIDADES = {'Abacaxi':'mil frutos','Coco-da-baía':'mil frutos'}

def norm(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)

def numero(v):
    if isinstance(v, (int, float)): return v
    return 0 if v == '-' else None  # '..', '...' e 'X' viram vazio

TEMPORARIAS = ['Abacaxi', 'Abóbora', 'Alfafa fenada', 'Alface', 'Algodão herbáceo (em caroço)', 'Alho', 'Amendoim (em casca)', 'Arroz (em casca)', 'Aveia (em grão)', 'Batata-doce', 'Batata-inglesa', 'Cana-de-açúcar', 'Cana para forragem', 'Canola', 'Cebola', 'Cenoura', 'Centeio (em grão)', 'Cevada (em grão)', 'Chuchu', 'Ervilha (em grão)', 'Fava (em grão)', 'Feijão (em grão)', 'Fumo (em folha)', 'Gergelim', 'Girassol (em grão)', 'Inhame', 'Juta (fibra)', 'Linho (semente)', 'Malva (fibra)', 'Mamona (baga)', 'Mandioca', 'Melancia', 'Melão', 'Milho (em grão)', 'Milho verde', 'Morango', 'Pimentão', 'Rami (fibra)', 'Repolho', 'Soja (em grão)', 'Sorgo (em grão)', 'Tomate', 'Trigo (em grão)', 'Triticale (em grão)']

def tabela5457(pasta, idx, ufs_com_tab3, linhas, faltando):
    """Lavouras temporárias por município a partir da Tabela 5457 (temporárias e permanentes juntas,
    um produto por coluna, uma variável por aba ou por arquivo). Usada só para UFs sem arquivo Tabela_3_*."""
    var_de = {'Área plantada': 'area', 'Quantidade produzida': 'qtd', 'Valor da produção': 'valor'}
    dados = defaultdict(dict)  # (uf, nome) -> produto -> {area, qtd, valor}
    for arq in sorted(glob.glob(os.path.join(pasta, '*.xlsx'))):
        if os.path.basename(arq).startswith('Tabela_'): continue
        wb = openpyxl.load_workbook(arq, read_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows or not isinstance(rows[0][0], str) or not rows[0][0].startswith('Tabela 5457') or rows[2][0] != 'Município': continue
            var = next((k for t, k in var_de.items() if t in (rows[1][0] or '')), None)
            if not var: continue
            cab = rows[4]
            for r in rows[5:]:
                mm = re.match(r'^(.*) \(([A-Z]{2})\)$', r[0] or '')
                if not mm: continue
                nome, uf = mm.group(1).strip(), mm.group(2)
                if uf in ufs_com_tab3: continue
                for h, v in zip(cab[1:], r[1:]):
                    if not h: continue
                    p = ROTULOS.get(h.replace('*', '').strip(), h.replace('*', '').strip())
                    if p in TEMPORARIAS: dados[(uf, nome)].setdefault(p, {})[var] = numero(v)
    for (uf, nome0), prods in dados.items():
        nome = CORRECOES.get((uf, nome0), nome0)
        m = idx.get((UFMAP[uf], norm(nome)))
        if not m: faltando.add((uf, nome)); continue
        ta = tv = None
        for p, d in prods.items():
            a, q, v = d.get('area'), d.get('qtd'), d.get('valor')
            if not a and not q and not v: continue
            linhas.append(dict(cod_ibge=m['codigo_ibge'], municipio=m['nome'], uf=uf, tipo_lavoura='Temporária', produto=p,
                               area_ha=a, quantidade=q, unidade_quantidade=UNIDADES.get(p, 't'), valor_mil_reais=v))
            if a is not None: ta = (ta or 0) + a
            if v is not None: tv = (tv or 0) + v
        if ta or tv:
            linhas.append(dict(cod_ibge=m['codigo_ibge'], municipio=m['nome'], uf=uf, tipo_lavoura='Temporária', produto='Total',
                               area_ha=ta, quantidade=None, unidade_quantidade=None, valor_mil_reais=tv))

def main(pasta):
    oficiais = list(csv.DictReader(open(os.path.join(RAIZ, 'scripts', 'municipios_ibge.csv'), encoding='utf-8')))
    idx = {(int(m['codigo_uf']), norm(m['nome'])): m for m in oficiais}
    linhas, faltando = [], set()
    for arq in sorted(glob.glob(os.path.join(pasta, 'Tabela_3_*.xlsx')) + glob.glob(os.path.join(pasta, 'Tabela_6_*.xlsx'))):
        tipo = 'Temporária' if os.path.basename(arq).startswith('Tabela_3_') else 'Permanente'
        wb = openpyxl.load_workbook(arq, read_only=True)
        for ws in wb.worksheets:
            if ws.title == 'Notas': continue
            dados = list(ws.iter_rows(values_only=True))
            produto = dados[2][0].split(' - ', 1)[1].replace('*', '').strip()
            produto = ROTULOS.get(produto, produto)
            uf = re.search(r' - ([A-Z]{2})$', dados[3][0]).group(1)
            for r in dados[5:]:
                if not isinstance(r[0], str) or len(r[0]) - len(r[0].lstrip(' ')) != 6: continue
                nome = CORRECOES.get((uf, r[0].strip()), r[0].strip())
                m = idx.get((UFMAP[uf], norm(nome)))
                if not m: faltando.add((uf, nome)); continue
                a, q, v = numero(r[1]), numero(r[3]), numero(r[5])
                if not a and not q and not v: continue
                linhas.append(dict(cod_ibge=m['codigo_ibge'], municipio=m['nome'], uf=uf, tipo_lavoura=tipo, produto=produto,
                                   area_ha=a, quantidade=None if produto == 'Total' else q,
                                   unidade_quantidade=None if produto == 'Total' else UNIDADES.get(produto, 't'), valor_mil_reais=v))
    ufs_tab3 = {re.search(r' - ([A-Z]{2})$', list(openpyxl.load_workbook(a, read_only=True).worksheets[0].iter_rows(min_row=4, max_row=4, values_only=True))[0][0]).group(1)
                for a in glob.glob(os.path.join(pasta, 'Tabela_3_*.xlsx'))}
    tabela5457(pasta, idx, ufs_tab3, linhas, faltando)
    if faltando:
        print('Municípios sem correspondência (inclua em CORRECOES):', sorted(faltando))
    soma = defaultdict(lambda: [None, None])
    for l in linhas:
        if l['produto'] == 'Total':
            s = soma[(l['cod_ibge'], l['municipio'], l['uf'])]
            if l['area_ha'] is not None: s[0] = (s[0] or 0) + l['area_ha']
            if l['valor_mil_reais'] is not None: s[1] = (s[1] or 0) + l['valor_mil_reais']
    for (cod, nome, uf), (a, v) in soma.items():
        linhas.append(dict(cod_ibge=cod, municipio=nome, uf=uf, tipo_lavoura='Todas', produto='Total', area_ha=a, quantidade=None, unidade_quantidade=None, valor_mil_reais=v))

    os.makedirs(os.path.join(RAIZ, 'dados'), exist_ok=True)
    cols = ['cod_ibge','municipio','uf','tipo_lavoura','produto','area_ha','quantidade','unidade_quantidade','valor_mil_reais']
    with open(os.path.join(RAIZ, 'dados', f'pam_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, cols); w.writeheader()
        for l in linhas: w.writerow({k: ('' if l[k] is None else l[k]) for k in cols})
    cods = sorted({l['cod_ibge'] for l in linhas}); mi = {c: i for i, c in enumerate(cods)}
    nomes = {l['cod_ibge']: (l['municipio'], l['uf']) for l in linhas}
    produtos, series = {}, {}
    for l in linhas:
        k = l['tipo_lavoura'] + '|' + l['produto']
        produtos.setdefault(k, {'chave': k, 'tipo': l['tipo_lavoura'], 'produto': l['produto'], 'unidade': l['unidade_quantidade']})
        series.setdefault(k, []).append([mi[l['cod_ibge']], l['area_ha'], l['quantidade'], l['valor_mil_reais']])
    saida = {'fonte': f'IBGE, Produção Agrícola Municipal {ANO} (Tabelas 1612 e 1613).',
             'colunas_series': ['municipio_idx', 'area_ha', 'quantidade', 'valor_mil_reais'],
             'municipios': [[c, *nomes[c]] for c in cods], 'produtos': list(produtos.values()), 'series': series}
    json.dump(saida, open(os.path.join(RAIZ, 'dados', f'pam_{ANO}.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'{len(linhas)} linhas, {len(cods)} municípios. Arquivos gravados em dados/.')

REBANHOS = {'Bovino':'Bovino','Bubalino':'Bubalino','Equino':'Equino','Suíno - total':'Suíno (total)','Suíno - matrizes de suínos':'Suíno – matrizes',
            'Caprino':'Caprino','Ovino':'Ovino','Galináceos - total':'Galináceos (total)','Galináceos - galinhas':'Galináceos – galinhas','Codornas':'Codornas'}

def pecuaria(pasta):
    oficiais = list(csv.DictReader(open(os.path.join(RAIZ, 'scripts', 'municipios_ibge.csv'), encoding='utf-8')))
    idx = {(int(m['codigo_uf']), norm(m['nome'])): m for m in oficiais}
    linhas, faltando = [], set()
    for arq in sorted(glob.glob(os.path.join(pasta, 'Tabela_2_*.xlsx'))):
        dados = list(openpyxl.load_workbook(arq, read_only=True).worksheets[0].iter_rows(values_only=True))
        uf = re.search(r' - ([A-Z]{2})$', dados[3][0]).group(1)
        cab = dados[4][1:11]
        for r in dados[5:]:
            if not isinstance(r[0], str) or len(r[0]) - len(r[0].lstrip(' ')) != 6: continue
            nome = CORRECOES.get((uf, r[0].strip()), r[0].strip())
            m = idx.get((UFMAP[uf], norm(nome)))
            if not m: faltando.add((uf, nome)); continue
            for h, v in zip(cab, r[1:11]):
                if isinstance(v, (int, float)) and v > 0:
                    linhas.append(dict(cod_ibge=m['codigo_ibge'], municipio=m['nome'], uf=uf, rebanho=REBANHOS[h], efetivo_cabecas=int(v)))
    if not linhas: return
    if faltando:
        print('Pecuária - municípios sem correspondência (inclua em CORRECOES):', sorted(faltando))
    cols = ['cod_ibge', 'municipio', 'uf', 'rebanho', 'efetivo_cabecas']
    with open(os.path.join(RAIZ, 'dados', f'ppm_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, cols); w.writeheader(); w.writerows(linhas)
    cods = sorted({l['cod_ibge'] for l in linhas}); mi = {c: i for i, c in enumerate(cods)}
    nomes = {l['cod_ibge']: (l['municipio'], l['uf']) for l in linhas}
    ordem = list(REBANHOS.values())
    series = {k: [] for k in ordem}
    for l in linhas: series[l['rebanho']].append([mi[l['cod_ibge']], l['efetivo_cabecas']])
    saida = {'fonte': f'IBGE, Pesquisa da Pecuária Municipal {ANO} (Tabela 3939), efetivo dos rebanhos em 31/12, em cabeças.',
             'municipios': [[c, *nomes[c]] for c in cods], 'rebanhos': ordem, 'series': series}
    json.dump(saida, open(os.path.join(RAIZ, 'dados', f'ppm_{ANO}.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'Pecuária: {len(linhas)} linhas, {len(cods)} municípios.')

# ---------- crédito rural (Banco Central, Matriz de Dados do Crédito Rural) ----------
PROGRAMAS = {'FINANCIAMENTO SEM VÍNCULO A PROGRAMA ESPECÍFICO': 'Sem vínculo a programa', 'PRONAF': 'Pronaf', 'PRONAMP': 'Pronamp', 'FUNCAFÉ': 'Funcafé',
             'MODERFROTA': 'Moderfrota', 'PCA': 'PCA (armazéns)', 'RENOVAGRO': 'RenovAgro', 'INOVAGRO': 'Inovagro', 'PROIRRIGA': 'Proirriga',
             'PRODECOOP': 'Prodecoop', 'PROCAP-AGRO': 'Procap-Agro', 'MODERAGRO': 'Moderagro', 'PROGRAMA NACIONAL DE CRÉDITO FUNDIÁRIO': 'Crédito Fundiário'}
LINHAS_PRONAF = {'CUSTEIO (MCR 10-4)': 'Custeio', 'MAIS ALIMENTOS (MCR 10-5)': 'Mais Alimentos', 'MICROCRÉDITO PRODUTIVO RURAL - GRUPO B (MCR 10-13)': 'Microcrédito (Grupo B)',
                 'CREDITO DE  INVESTIMENTO - PRONAF BIOECONOMIA (MCR 10-16)': 'Bioeconomia', 'PRONAF INDUSTRIALIZAÇÃO AGROIND. FAMILIAR (MCR10-11)': 'Industrialização da agroindústria familiar',
                 'REFORMA AGRÁRIA BENEFICIÁRIOS PNCF, PNRA, PCRF (MCR10-3)': 'Reforma agrária', 'COTAS PARTES (MCR 10-12)': 'Cotas-partes', 'AGROINDÚSTRIA (INVESTIMENTO) (MCR 10-6)': 'Agroindústria',
                 'MULHER (MCR 10-9)': 'Mulher', 'CRÉD INVEST CONVIV SEMIÁRIDO (PRONAF SEMIÁRIDO) MCR 10-8': 'Semiárido',
                 'CRÉDITO INVEST SIST AGROFLORESTAIS (PRONAF FLORESTA) MCR10-7': 'Floresta (sistemas agroflorestais)', 'CRÉD INVEST AGROECOLOGIA (PRONAF AGROECOLOGIA)MCR 10-14': 'Agroecologia', 'JOVEM (MCR 10-10)': 'Jovem'}
SEGMENTOS = {'BANCO PÚBLICO': 'Banco público', 'BANCO PRIVADO': 'Banco privado', 'COOPERATIVA DE CRÉDITO': 'Cooperativa de crédito', 'BANCO COOPERATIVO': 'Banco cooperativo',
             'BANCO DE DESENVOLVIMENTO E AGÊNCIA DE FOMENTO': 'Banco de desenvolvimento e agência de fomento', 'SOCIEDADE DE CRÉDITO, FINANCIAMENTO E INVESTIMENTO': 'Financeira (SCFI)'}
FINALIDADES = [('CUSTEIO', 'Custeio'), ('INVESTIMENTO', 'Investimento'), ('COMERCIALIZACAO', 'Comercialização'), ('INDUSTRIALIZACAO', 'Industrialização')]

def credito(pasta):
    """Agrega a Matriz de Dados do Crédito Rural (arquivo MATRIZ_CR_REGIAO*.csv, separador ';', decimal ',')
    por município, programa, linha, setor, finalidade e instituição financeira."""
    arqs = sorted(glob.glob(os.path.join(pasta, 'MATRIZ_CR_REGIAO*.csv')))
    if not arqs: return
    oficiais = {m['codigo_ibge']: m for m in csv.DictReader(open(os.path.join(RAIZ, 'scripts', 'municipios_ibge.csv'), encoding='utf-8'))}
    num = lambda x: float(x.replace('.', '').replace(',', '.')) if x else 0.0
    acc, desconhecidos, nomes_if = defaultdict(lambda: [0.0, 0]), set(), {}
    valor_if = defaultdict(float)
    for arq in arqs:
        for r in csv.DictReader(open(arq, encoding='utf-8-sig'), delimiter=';'):
            chave_prog = re.split(r' - | \(', r['PROGRAMA'])[0].strip()
            prog = PROGRAMAS.get(chave_prog)
            if prog is None: desconhecidos.add(r['PROGRAMA']); continue
            sub = r['SUBPROGRAMA']
            linha = LINHAS_PRONAF.get(sub, sub.title()) if prog == 'Pronaf' else (sub.replace('RENOVAGRO ', '').capitalize() if prog == 'RenovAgro' else '')
            setor = 'Agrícola' if r['ATIVIDADE'].startswith('AGR') else 'Pecuária'
            cnpj = r['CNPJ_IF'].zfill(8)
            nomes_if[cnpj] = (r['NOME_IF'], SEGMENTOS.get(r['SEGMENTO_IF'], r['SEGMENTO_IF'].capitalize()))
            for col, fin in FINALIDADES:
                v, n = num(r[col + '_VALOR']), int(num(r[col + '_QTDE']))
                if v or n:
                    a = acc[(r['COD_MUNICIPIO_IBGE'], prog, linha, setor, fin, cnpj)]
                    a[0] += v; a[1] += n; valor_if[cnpj] += v
    if desconhecidos: print('Programas não reconhecidos (inclua em PROGRAMAS):', sorted(desconhecidos))
    chaves = sorted(acc)
    ordem_if = sorted(nomes_if, key=lambda c: -valor_if[c])
    with open(os.path.join(RAIZ, 'dados', 'instituicoes_financeiras.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['cnpj_if', 'nome', 'segmento'])
        for c in ordem_if: w.writerow([c, *nomes_if[c]])
    with open(os.path.join(RAIZ, 'dados', f'credito_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['cod_ibge', 'municipio', 'uf', 'programa', 'linha', 'atividade', 'finalidade', 'cnpj_if', 'valor_reais', 'contratos'])  # nome e tipo da instituição: instituicoes_financeiras.csv
        for k in chaves:
            m = oficiais.get(k[0], {'nome': '', 'codigo_uf': k[0][:2]})
            w.writerow([k[0], m['nome'], UFSIG[int(m['codigo_uf'])], *k[1:5], k[5], round(acc[k][0], 2), acc[k][1]])
    programas = [p for p in PROGRAMAS.values() if any(k[1] == p for k in chaves)]
    soma_linha = defaultdict(float)
    for k in chaves:
        if k[2]: soma_linha[k[2]] += acc[k][0]
    linhas = sorted(soma_linha, key=lambda l: -soma_linha[l])
    tipos = [s for s in SEGMENTOS.values() if any(nomes_if[c][1] == s for c in ordem_if)]
    fins = [f for _, f in FINALIDADES]
    ii = {c: i for i, c in enumerate(ordem_if)}
    cods = sorted({k[0] for k in chaves}); mi = {c: i for i, c in enumerate(cods)}
    dados = [[mi[k[0]], programas.index(k[1]), linhas.index(k[2]) if k[2] else -1, 0 if k[3] == 'Agrícola' else 1, fins.index(k[4]), ii[k[5]], round(acc[k][0]), acc[k][1]] for k in chaves]
    municipios = [[c, oficiais.get(c, {}).get('nome', c), UFSIG[int(c[:2])]] for c in cods]
    saida = {'fonte': f'Banco Central do Brasil, Matriz de Dados do Crédito Rural (SICOR), operações contratadas em {ANO} (data de emissão). Valores nominais em reais.',
             'municipios': municipios, 'programas': programas, 'linhas': linhas, 'atividades': ['Agrícola', 'Pecuária'], 'finalidades': fins, 'tipos_instituicao': tipos,
             'instituicoes': [[c, nomes_if[c][0], tipos.index(nomes_if[c][1])] for c in ordem_if],
             'colunas': ['municipio_idx', 'programa_idx', 'linha_idx (-1 = sem linha)', 'atividade_idx', 'finalidade_idx', 'instituicao_idx', 'valor_reais', 'contratos'], 'dados': dados}
    json.dump(saida, open(os.path.join(RAIZ, 'dados', f'credito_{ANO}.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'Crédito rural: {len(chaves)} linhas, {len(cods)} municípios, R$ {sum(a[0] for a in acc.values()) / 1e9:.1f} bi.')

def produtos(pasta):
    """Agrega a Matriz de Dados do Crédito Rural por produto, base por estado (arquivo MATRIZ_CR_PRODUTO_REGIAO*.csv),
    por estado, modalidade, produto, programa, linha, setor, finalidade e instituição. A base por município
    não é usada: nela o Banco Central omite o município de combinações com até 3 contratos (sigilo bancário)."""
    arqs = sorted(glob.glob(os.path.join(pasta, 'MATRIZ_CR_PRODUTO_REGIAO*.csv')))
    if not arqs: return
    nome_uf = {'ACRE': 'AC', 'ALAGOAS': 'AL', 'AMAPÁ': 'AP', 'AMAZONAS': 'AM', 'BAHIA': 'BA', 'CEARÁ': 'CE', 'DISTRITO FEDERAL': 'DF', 'ESPÍRITO SANTO': 'ES', 'GOIÁS': 'GO',
               'MARANHÃO': 'MA', 'MATO GROSSO': 'MT', 'MATO GROSSO DO SUL': 'MS', 'MINAS GERAIS': 'MG', 'PARÁ': 'PA', 'PARAÍBA': 'PB', 'PARANÁ': 'PR', 'PERNAMBUCO': 'PE',
               'PIAUÍ': 'PI', 'RIO DE JANEIRO': 'RJ', 'RIO GRANDE DO NORTE': 'RN', 'RIO GRANDE DO SUL': 'RS', 'RONDÔNIA': 'RO', 'RORAIMA': 'RR', 'SANTA CATARINA': 'SC',
               'SÃO PAULO': 'SP', 'SERGIPE': 'SE', 'TOCANTINS': 'TO'}
    num = lambda x: float(x.replace('.', '').replace(',', '.')) if x else 0.0
    cap = lambda t: t[:1].upper() + t[1:].lower()
    acc, desconhecidos = defaultdict(lambda: [0.0, 0, 0.0]), set()
    for arq in arqs:
        for r in csv.DictReader(open(arq, encoding='utf-8-sig'), delimiter=';'):
            prog = PROGRAMAS.get(re.split(r' - | \(', r['PROGRAMA'])[0].strip())
            if prog is None: desconhecidos.add(r['PROGRAMA']); continue
            sub = r['SUBPROGRAMA']
            linha = LINHAS_PRONAF.get(sub, sub.title()) if prog == 'Pronaf' else (sub.replace('RENOVAGRO ', '').capitalize() if prog == 'RenovAgro' else '')
            setor = 'Agrícola' if r['ATIVIDADE'].startswith('AGR') else 'Pecuária'
            for col, fin in FINALIDADES:
                v, n = num(r[col + '_VALOR']), int(num(r[col + '_QTDE']))
                area = num(r.get(col + '_AREA', '')) if col in ('CUSTEIO', 'INVESTIMENTO') else 0.0
                if v or n:
                    a = acc[(nome_uf[r['UF']], cap(r['MODALIDADE']), cap(r['PRODUTO']), prog, linha, setor, fin, r['CNPJ_IF'].zfill(8))]
                    a[0] += v; a[1] += n; a[2] += area
    if desconhecidos: print('Produtos - programas não reconhecidos (inclua em PROGRAMAS):', sorted(desconhecidos))
    chaves = sorted(acc)
    with open(os.path.join(RAIZ, 'dados', f'credito_produtos_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['uf', 'modalidade', 'produto', 'programa', 'linha', 'atividade', 'finalidade', 'cnpj_if', 'valor_reais', 'contratos', 'area_ha'])
        for k in chaves: w.writerow([*k, round(acc[k][0], 2), acc[k][1], round(acc[k][2], 2)])
    def ordem(i):
        s = defaultdict(float)
        for k in chaves: s[k[i]] += acc[k][0]
        return sorted(s, key=lambda x: -s[x])
    ufs = sorted({k[0] for k in chaves}); mods, prods, progs, ifs = ordem(1), ordem(2), ordem(3), ordem(7)
    lins = [l for l in ordem(4) if l]; fins = [f for _, f in FINALIDADES]
    ix = lambda L: {v: i for i, v in enumerate(L)}
    U, M, P, G, L, F, I = map(ix, [ufs, mods, prods, progs, lins, fins, ifs])
    dados = [[U[k[0]], M[k[1]], P[k[2]], G[k[3]], L.get(k[4], -1), 0 if k[5] == 'Agrícola' else 1, F[k[6]], I[k[7]], round(acc[k][0]), acc[k][1], round(acc[k][2], 1)] for k in chaves]
    saida = {'fonte': f'Banco Central do Brasil, Matriz de Dados do Crédito Rural por produto, base por estado, operações contratadas em {ANO}. Valores nominais em reais; área em hectares (custeio e investimento).',
             'ufs': ufs, 'modalidades': mods, 'produtos': prods, 'programas': progs, 'linhas': lins, 'atividades': ['Agrícola', 'Pecuária'], 'finalidades': fins, 'instituicoes_cnpj': ifs,
             'colunas': ['uf_idx', 'modalidade_idx', 'produto_idx', 'programa_idx', 'linha_idx (-1 = sem linha)', 'atividade_idx', 'finalidade_idx', 'instituicao_idx', 'valor_reais', 'contratos', 'area_ha'], 'dados': dados}
    json.dump(saida, open(os.path.join(RAIZ, 'dados', f'credito_produtos_{ANO}.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'Produtos financiados: {len(chaves)} linhas, R$ {sum(a[0] for a in acc.values()) / 1e9:.1f} bi.')

def pecuaria_estados(pasta):
    """Produção de origem animal (Tabela 74) e aquicultura (Tabela 3940) da PPM, por estado.
    Usa qualquer .xlsx da pasta cuja primeira planilha comece com "Tabela 74" ou "Tabela 3940"."""
    nomes_uf = {v: k for k, v in {'RO': 'Rondônia', 'AC': 'Acre', 'AM': 'Amazonas', 'RR': 'Roraima', 'PA': 'Pará', 'AP': 'Amapá', 'TO': 'Tocantins', 'MA': 'Maranhão',
                'PI': 'Piauí', 'CE': 'Ceará', 'RN': 'Rio Grande do Norte', 'PB': 'Paraíba', 'PE': 'Pernambuco', 'AL': 'Alagoas', 'SE': 'Sergipe', 'BA': 'Bahia',
                'MG': 'Minas Gerais', 'ES': 'Espírito Santo', 'RJ': 'Rio de Janeiro', 'SP': 'São Paulo', 'PR': 'Paraná', 'SC': 'Santa Catarina', 'RS': 'Rio Grande do Sul',
                'MS': 'Mato Grosso do Sul', 'MT': 'Mato Grosso', 'GO': 'Goiás', 'DF': 'Distrito Federal'}.items()}
    unidades = {'Quilogramas': 'kg', 'Mil litros': 'mil litros', 'Mil dúzias': 'mil dúzias', 'Milheiros': 'milheiros'}
    renomear = {'Peixes': 'Peixes (total)', 'Outros produtos (rã, jacaré, siri, caranguejo, lagosta, etc)': 'Outros produtos (rã, jacaré, siri, caranguejo, lagosta etc.)'}
    saida = []
    for arq in sorted(glob.glob(os.path.join(pasta, '*.xlsx'))):
        wb = openpyxl.load_workbook(arq, read_only=True)
        primeira = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True))[0] or ''
        grupo = 'Aquicultura' if primeira.startswith('Tabela 3940') else 'Origem animal' if primeira.startswith('Tabela 74 ') else None
        if not grupo: continue
        for ws in wb.worksheets:
            if ws.title == 'Notas': continue
            linhas = list(ws.iter_rows(values_only=True))
            prod = linhas[1][0].split(' - ', 1)[1].strip()
            if prod == 'Total': continue
            cab = linhas[4][1] or ''
            unidade = unidades.get(cab[cab.find('(') + 1:cab.rfind(')')], '') if '(' in cab else ''
            for r in linhas[5:]:
                uf = nomes_uf.get((r[0] or '').strip())
                if not uf: continue
                q, v = numero(r[1]), numero(r[2])
                if not q and not v: continue
                saida.append({'uf': uf, 'grupo': grupo, 'produto': renomear.get(prod, prod), 'agregado': 1 if prod == 'Peixes' else 0, 'unidade': unidade,
                              'quantidade': '' if q is None or not unidade else q, 'valor_mil_reais': '' if v is None else v})
    if not saida: return
    with open(os.path.join(RAIZ, 'dados', f'pecuaria_produtos_uf_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, list(saida[0])); w.writeheader(); w.writerows(saida)
    print(f'Origem animal e aquicultura (por estado): {len(saida)} linhas.')

def car(pasta):
    """Cadastro Ambiental Rural (Painel de Regularização Ambiental, MMA): número de cadastros e área cadastrada por
    município, tamanho, tipo de imóvel e condição do CAR. Usa duas planilhas exportadas do Painel, uma com a coluna
    "n total de cadastros" e outra com "Área (ha) total cadastrada", cada uma com uma aba por faixa de módulo fiscal."""
    planilhas = {}
    for arq in sorted(glob.glob(os.path.join(pasta, '*.xlsx'))):
        wb = openpyxl.load_workbook(arq, read_only=True)
        cab = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True))
        if len(cab) >= 7 and cab[0] == 'UF' and cab[3] == 'Módulo Fiscal':
            planilhas['n' if 'cadastros' in str(cab[6]) else 'area'] = wb
    if set(planilhas) != {'n', 'area'}: return
    oficiais = list(csv.DictReader(open(os.path.join(RAIZ, 'scripts', 'municipios_ibge.csv'), encoding='utf-8')))
    idx = {(int(m['codigo_uf']), norm(m['nome'])): m for m in oficiais}
    tamanho = {'0 a 4': 'Pequena', '4 a 15': 'Média', 'acima de 15': 'Grande'}
    acc, faltando = defaultdict(lambda: [0, 0.0]), set()
    for chave, wb in planilhas.items():
        for ws in wb.worksheets:
            for r in list(ws.iter_rows(values_only=True))[1:]:
                if not r[0] or r[0] == 'Total' or not r[1]: continue
                m = idx.get((UFMAP[r[0]], norm(CORRECOES.get((r[0], r[1]), r[1]))))
                if not m: faltando.add((r[0], r[1])); continue
                a = acc[(m['codigo_ibge'], tamanho[r[3]], r[4], r[5])]
                if chave == 'n': a[0] += int(r[6] or 0)
                else: a[1] += float(r[6] or 0)
    if faltando: print('CAR - municípios sem correspondência (inclua em CORRECOES):', sorted(faltando))
    chaves = sorted(k for k in acc if acc[k][0] or acc[k][1])
    nomes = {m['codigo_ibge']: m for m in oficiais}
    with open(os.path.join(RAIZ, 'dados', f'car_{ANO}.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['cod_ibge', 'municipio', 'uf', 'tamanho', 'tipo', 'condicao', 'cadastros', 'area_ha'])
        for k in chaves: w.writerow([k[0], nomes[k[0]]['nome'], UFSIG[int(k[0][:2])], *k[1:], acc[k][0], round(acc[k][1], 2)])
    cods = sorted({k[0] for k in chaves}); mi = {c: i for i, c in enumerate(cods)}
    tams, tipos = ['Pequena', 'Média', 'Grande'], ['IRU', 'AST', 'PCT']
    soma = defaultdict(int)
    for k in chaves: soma[k[3]] += acc[k][0]
    conds = sorted(soma, key=lambda c: -soma[c])
    saida = {'fonte': 'Ministério do Meio Ambiente e Mudança do Clima, Painel de Regularização Ambiental (CAR). Número de cadastros e área cadastrada (ha), declarada.',
             'municipios': [[c, nomes[c]['nome'], UFSIG[int(c[:2])]] for c in cods], 'tamanhos': tams, 'tipos': tipos, 'condicoes': conds,
             'colunas': ['municipio_idx', 'tamanho_idx', 'tipo_idx', 'condicao_idx', 'cadastros', 'area_ha'],
             'dados': [[mi[k[0]], tams.index(k[1]), tipos.index(k[2]), conds.index(k[3]), acc[k][0], round(acc[k][1], 1)] for k in chaves]}
    json.dump(saida, open(os.path.join(RAIZ, 'dados', f'car_{ANO}.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'CAR: {len(chaves)} linhas, {sum(a[0] for a in acc.values())} cadastros, {sum(a[1] for a in acc.values()) / 1e6:.1f} milhões de ha.')

if __name__ == '__main__':
    if len(sys.argv) != 2: sys.exit(__doc__)
    if glob.glob(os.path.join(sys.argv[1], 'Tabela_[36]_*.xlsx')): main(sys.argv[1])
    pecuaria(sys.argv[1])
    credito(sys.argv[1])
    produtos(sys.argv[1])
    pecuaria_estados(sys.argv[1])
    car(sys.argv[1])
